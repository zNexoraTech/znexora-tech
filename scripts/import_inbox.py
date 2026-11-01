"""Turns Word (.docx) and PowerPoint (.pptx) files you upload into website articles.

How to use: upload your files into the inbox folder on GitHub, inside a folder named after
the section, for example:

    inbox/theory/Zero Trust Explained.docx
    inbox/theory/Zero Trust Explained.pptx      (optional, same name = same article)
    inbox/practical/Harden SSH.docx

The robot converts them and opens a draft (pull request) for your approval, exactly like the
other robots. When you merge it, the article is published and shared everywhere.

Word file    -> the article text, headings, lists, tables and pictures
PowerPoint   -> a PDF download, a slide gallery, and the first slide as the cover picture
"""
import argparse
import hashlib
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from common import (PR_FOOTER, ROOT, automation, commit_state, dump_front_matter, git, load_json,
                    save_json, slugify, today)

INBOX = ROOT / "inbox"
SECTIONS = {"news", "weekly-recap", "exploited-vulnerabilities", "theory", "practical", "courses", "cheatsheets"}
ALIASES = {"articles": "theory", "article": "theory", "training": "practical", "course": "courses",
           "cheat-sheets": "cheatsheets", "cheatsheet": "cheatsheets"}  # folder name -> section
PICTURES = (".png", ".jpg", ".jpeg", ".webp")
IMAGES = ROOT / "assets" / "images"
FILES = ROOT / "assets" / "files"


def file_hash(path):
    return hashlib.sha1(path.read_bytes()).hexdigest()


def groups():
    """Pair .docx and .pptx files that share a name. Returns {(section, name): {"docx": p, "pptx": p}}."""
    found = {}
    for path in sorted(INBOX.rglob("*")):
        if path.suffix.lower() not in (".docx", ".pptx") + PICTURES or path.name.startswith("~$"):
            continue
        section = path.parent.name.lower() if path.parent != INBOX else "theory"
        section = ALIASES.get(section, section)
        if section not in SECTIONS:
            section = "theory"
        kind = "picture" if path.suffix.lower() in PICTURES else path.suffix.lower()[1:]
        found.setdefault((section, path.stem.strip()), {})[kind] = path
    return found


# ---------------------------------------------------------------- Word
def convert_docx(path, slug):
    import mammoth
    from markdownify import markdownify

    folder = IMAGES / slug
    count = [0]

    def save_image(image):
        count[0] += 1
        ext = {"image/png": "png", "image/jpeg": "jpg", "image/gif": "gif", "image/webp": "webp"}.get(image.content_type)
        if not ext:  # unusual formats (e.g. EMF drawings) are skipped
            return {"src": ""}
        folder.mkdir(parents=True, exist_ok=True)
        name = f"img-{count[0]}.{ext}"
        with image.open() as data:
            (folder / name).write_bytes(data.read())
        return {"src": f"../../assets/images/{slug}/{name}", "alt": image.alt_text or ""}

    with path.open("rb") as f:
        result = mammoth.convert_to_html(f, convert_image=mammoth.images.img_element(save_image))
    html = re.sub(r'<img src=""[^>]*>', "", result.value)
    md = markdownify(html, heading_style="ATX", bullets="-", strip=["span"], table_infer_header=True)
    md = re.sub(r"\n{3,}", "\n\n", md).strip()

    # Title: the first big heading, or the file name
    title = path.stem.strip()
    m = re.match(r"#\s+(.+)\n", md + "\n")
    if m:
        title = m.group(1).strip()
        md = md[m.end():].lstrip()
    md = re.sub(r"^#\s+", "## ", md, flags=re.M)  # article headings start at level 2 on the website
    first = next((p for p in md.split("\n\n") if p.strip() and not p.startswith(("#", "!", "-", "|", "*"))), "")
    summary = re.sub(r"[*_`\[\]]|\(http[^)]*\)", "", first).strip()
    if first and len(summary) <= 220:  # shown in the Key points box, so don't repeat it as the opening
        md = md.replace(first, "", 1).lstrip()
    summary = summary[:220].rsplit(" ", 1)[0] + "…" if len(summary) > 220 else summary
    return title, summary, md


# ---------------------------------------------------------------- PowerPoint
def _shapes(shapes):
    """All shapes on a slide, including those inside groups."""
    for shape in shapes:
        if shape.shape_type == 6:  # group
            yield from _shapes(shape.shapes)
        else:
            yield shape


def read_slides(path):
    """Titles, bullet points and photos from each slide. Speaker notes are left out on purpose."""
    from pptx import Presentation
    deck = Presentation(str(path))
    slides, counts = [], {}
    for slide in deck.slides:
        title, lines, pics = "", [], []
        for shape in _shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                ptype = str(shape.placeholder_format.type) if shape.is_placeholder else ""
                is_title = "TITLE" in ptype and "SUB" not in ptype
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if not text:
                        continue
                    if is_title and not title:
                        title = text
                    else:
                        lines.append((para.level, text))
            image = getattr(shape, "image", None) if shape.shape_type in (13, 14) or hasattr(shape, "image") else None
            try:
                if image is not None:
                    blob = image.blob
                    key = hashlib.sha1(blob).hexdigest()
                    counts[key] = counts.get(key, 0) + 1
                    pics.append((key, blob))
            except Exception:
                pass
        slides.append({"title": title, "lines": lines, "pics": pics})
    for sl in slides:  # drop logos repeated on many slides
        sl["pics"] = [(k, b) for k, b in sl["pics"] if counts.get(k, 0) < 3]
    return slides


def save_picture(blob, dest_base):
    """Save an embedded photo as PNG/JPG for the web. Returns the file name, or "" if unusable."""
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(blob))
        im.load()
    except Exception:
        return ""
    if im.width < 160 or im.height < 160:  # icons and bullets, not photos
        return ""
    ext = "jpg" if im.format == "JPEG" else "png"
    if im.mode not in ("RGB", "RGBA", "L", "P"):
        im = im.convert("RGBA")
    if ext == "jpg" and im.mode != "RGB":
        im = im.convert("RGB")
    im.thumbnail((1600, 1600))
    dest = dest_base.with_suffix("." + ext)
    dest.parent.mkdir(parents=True, exist_ok=True)
    im.save(dest, quality=85) if ext == "jpg" else im.save(dest, optimize=True)
    return dest.name


def soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def render_slides(path, slug):
    """Turn every slide into a picture and the deck into a PDF. Returns (has_pdf, [slide picture names])."""
    if not soffice():
        print("  LibreOffice is not installed, so slide pictures are skipped (photos are taken from the file instead).")
        return False, []
    work = Path(tempfile.mkdtemp())
    try:
        subprocess.run([soffice(), "--headless", "--convert-to", "pdf", "--outdir", str(work), str(path)],
                       check=True, capture_output=True, timeout=300)
        pdf = next(work.glob("*.pdf"))
        FILES.mkdir(parents=True, exist_ok=True)
        shutil.copy(pdf, FILES / f"{slug}.pdf")
        subprocess.run(["pdftoppm", "-jpeg", "-jpegopt", "quality=82", "-scale-to", "1280", str(pdf), str(work / "s")],
                       check=True, capture_output=True, timeout=300)
        rendered = sorted(work.glob("s-*.jpg"), key=lambda p: int(re.findall(r"\d+", p.stem)[-1]))
        folder = IMAGES / slug / "slides"
        folder.mkdir(parents=True, exist_ok=True)
        names = []
        for i, sl in enumerate(rendered, 1):
            shutil.copy(sl, folder / f"slide-{i}.jpg")
            names.append(f"slide-{i}.jpg")
        print(f"  {len(names)} slide picture(s) and a PDF created.")
        return True, names
    except Exception as exc:
        print(f"  Slide conversion failed ({exc}); using the photos inside the file instead.")
        return False, []
    finally:
        shutil.rmtree(work, ignore_errors=True)


def make_cover_from(src, slug):
    from PIL import Image
    im = Image.open(src).convert("RGB")
    ratio = 1200 / 630
    w, h = im.size
    if w / h > ratio:
        nw = int(h * ratio); im = im.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
    else:
        nh = int(w / ratio); top = 0 if h > w else (h - nh) // 2; im = im.crop((0, top, w, top + nh))
    im.resize((1200, 630), Image.LANCZOS).save(IMAGES / f"{slug}.jpg", quality=85, optimize=True)


def pptx_article(path, slug, with_text):
    """Markdown for a deck. with_text=True (slides only): one section per slide with its picture and points.
    with_text=False (a Word file has the text): a slide gallery. Returns (title, summary, markdown)."""
    slides = read_slides(path)
    has_pdf, rendered = render_slides(path, slug)
    pic_dir = IMAGES / slug / "photos"

    def pictures_for(i, sl):
        if i <= len(rendered):
            return [f"../../assets/images/{slug}/slides/{rendered[i - 1]}"]
        names = [save_picture(blob, pic_dir / f"slide{i}-{n}") for n, (_, blob) in enumerate(sl["pics"], 1)]
        return [f"../../assets/images/{slug}/photos/{n}" for n in names if n]

    # cover: slide 1 picture, or the first photo in the deck
    if rendered:
        make_cover_from(IMAGES / slug / "slides" / rendered[0], slug)
    else:
        for i, sl in enumerate(slides, 1):
            pics = pictures_for(i, sl)
            if pics:
                make_cover_from(IMAGES / pics[0].split("/assets/images/")[1], slug)
                break

    title = slides[0]["title"] if slides else ""
    first_text = " ".join(t for _, t in slides[0]["lines"]) if slides else ""
    summary = first_text[:220]
    parts = []
    if has_pdf:
        parts.append(f"[Download the slides as PDF](../../assets/files/{slug}.pdf)")
    if with_text:
        for i, sl in enumerate(slides, 1):
            if i == 1:
                continue  # the title slide is the article's title, summary and cover
            block = [f"## {sl['title'] or f'Slide {i}'}"]
            block += [f"![{sl['title'] or f'Slide {i}'}]({src})" for src in pictures_for(i, sl)]
            if sl["lines"]:
                block.append("\n".join(f"{'    ' * min(level, 2)}- {text}" for level, text in sl["lines"]))
            parts.append("\n\n".join(block))
    else:
        gallery = [f"![Slide {i}]({src})" for i, sl in enumerate(slides, 1) for src in pictures_for(i, sl)]
        if gallery:
            parts.insert(0, "## Slides")
            parts += gallery
    if not summary and len(slides) > 1:
        summary = " ".join(t for _, t in slides[1]["lines"])[:220]
    return title, summary, "\n\n".join(parts)


# ---------------------------------------------------------------- drafts
def make_article(section, name, files, site):
    docx, pptx, picture = files.get("docx"), files.get("pptx"), files.get("picture")
    clean = re.sub(r"[-_]+", " ", name).strip()
    title, summary, body = (clean.title() if clean.islower() else clean, "", "")
    slug_hint = slugify(name)
    if docx:
        title, summary, body = convert_docx(docx, slug_hint)
    slug = slugify(title)
    if slug != slug_hint and (IMAGES / slug_hint).exists():  # keep pictures under the final name
        (IMAGES / slug_hint).rename(IMAGES / slug)
        body = body.replace(f"/images/{slug_hint}/", f"/images/{slug}/")
    slides_md = ""
    if pptx and not docx:  # slides only: title, summary and sections come from the slides
        first = read_slides(pptx)[:1]
        if first and first[0]["title"]:
            title, slug = first[0]["title"], slugify(first[0]["title"])
        _, summary, body = pptx_article(pptx, slug, with_text=True)
    elif pptx:  # a Word file has the text; the slides become a gallery
        _, _, slides_md = pptx_article(pptx, slug, with_text=False)
    if picture:  # a cheat sheet or infographic: shown in full, downloadable, and used as the cover
        IMAGES.mkdir(parents=True, exist_ok=True)
        dest = IMAGES / f"{slug}{picture.suffix.lower()}"
        shutil.copy(picture, dest)
        rel = f"../../assets/images/{dest.name}"
        kind = "cheat sheet" if section == "cheatsheets" else "picture"
        body = (f"![{title}]({rel})\n\n[Download the {kind} ({picture.suffix[1:].upper()})]({rel})"
                + (f"\n\n{body}" if body else ""))
        if not summary:
            what = title if "cheat sheet" in title.lower() or "cheatsheet" in title.lower() else f"{title} cheat sheet"
            summary = (f"A free one-page {what} from ZNexora Tech. Download it, print it, and share it with your team."
                       if section == "cheatsheets" else f"{title}, from ZNexora Tech.")
    if not summary:
        summary = f"{title}: slides and key points from ZNexora Tech."
    meta = {"title": title, "date": today().isoformat(), "category": section, "summary": summary,
            "author": site.get("author") or site["name"], "drafted_by": "upload"}
    full = "\n\n".join(x for x in (body, slides_md) if x)
    return slug, title, dump_front_matter(meta, full or "Write your article here.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    site, _, _ = automation("articles")
    done = load_json("imported.json", {})
    # forget files that are no longer in the inbox (already published), so the same file can be uploaded again
    done = {k: v for k, v in done.items() if (ROOT / k).exists()}
    todo = {k: v for k, v in groups().items()
            if any(done.get(str(p.relative_to(ROOT))) != file_hash(p) for p in v.values())}
    if not todo:
        return print("Nothing new in the inbox.")
    for (section, name), files in todo.items():
        print(f"- {section}: {name} ({', '.join(sorted(files))})")
        sources = [str(p.relative_to(ROOT)) for p in files.values()]
        hashes = {s: file_hash(ROOT / s) for s in sources}
        if args.dry_run:
            slug, title, content = make_article(section, name, files, site)
            (ROOT / "content" / "posts" / f"{slug}.md").write_text(content, encoding="utf-8")
            print(f"  wrote content/posts/{slug}.md")
            done.update(hashes)
            continue
        try:
            branch_slug = slugify(name)
            git("checkout", "-B", f"draft/upload-{branch_slug}", "main")
            slug, title, content = make_article(section, name, files, site)
            post = ROOT / "content" / "posts" / f"{slug}.md"
            post.write_text(content, encoding="utf-8")
            git("add", str(post))
            for extra in [IMAGES / slug, IMAGES / f"{slug}.jpg", FILES / f"{slug}.pdf"] + [IMAGES / f"{slug}{x}" for x in PICTURES]:
                if extra.exists():
                    git("add", str(extra))
            git("rm", "-q", *sources)  # the inbox empties itself when you approve
            git("commit", "-m", f"Draft from upload: {title}")
            git("push", "--force", "origin", f"draft/upload-{branch_slug}")
            body = (f"Article converted from your upload ({', '.join(Path(s).name for s in sources)}).\n\n"
                    f"- [ ] Title, headings and pictures look right (see **Files changed**)\n"
                    f"- [ ] The summary in the top section reads well\n\n"
                    f"When you merge, it is published on the website and shared to your channels.\n{PR_FOOTER}")
            subprocess.run(["gh", "pr", "create", "--base", "main", "--head", f"draft/upload-{branch_slug}",
                            "--title", f"Draft: {title}", "--body", body], cwd=ROOT, check=True,
                           capture_output=True, text=True)
            print(f"  opened draft: {title}")
            done.update(hashes)
        except subprocess.CalledProcessError as exc:
            print(f"  could not create the draft: {(exc.stderr or str(exc)).strip()[:300]}")
        finally:
            subprocess.run(["git", "checkout", "-f", "main"], cwd=ROOT, capture_output=True)
            subprocess.run(["git", "clean", "-fdq", "assets/images", "assets/files", "content/posts"], cwd=ROOT,
                           capture_output=True)
    save_json("imported.json", done)
    commit_state("Robot memory: imported uploads", args.dry_run)


if __name__ == "__main__":
    main()
